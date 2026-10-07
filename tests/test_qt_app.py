"""The Qt application: exports, both analysis pages, undo/redo, sessions, reports (offscreen)."""
import csv
import os
import time
from pathlib import Path

import numpy as np
from PySide6.QtCore import Qt
import pytest

pytest.importorskip("PySide6")

import exports  # noqa: E402
import peak_fitting  # noqa: E402
from selftest import write_ftir_csv, write_xrd_xy  # noqa: E402
from trace_model import Trace  # noqa: E402


@pytest.fixture(scope="session")
def qapp():
    from PySide6.QtWidgets import QApplication
    app = QApplication.instance() or QApplication([])
    app.setStyle("Fusion")
    return app


@pytest.fixture
def ftir_csv(tmp_path):
    return str(write_ftir_csv(tmp_path / "sample_ftir.csv"))


@pytest.fixture
def xrd_xy(tmp_path):
    return str(write_xrd_xy(tmp_path / "sample_xrd.xy"))


@pytest.fixture
def window(qapp, tmp_path, monkeypatch):
    from PySide6.QtCore import QSettings
    import app_config
    import ui_common

    QSettings.setDefaultFormat(QSettings.IniFormat)
    QSettings.setPath(QSettings.IniFormat, QSettings.UserScope, str(tmp_path / "settings"))
    monkeypatch.setattr(app_config, "get_config_dir", lambda: str(tmp_path))
    messages = []
    monkeypatch.setattr(ui_common, "show_message",
                        lambda parent, title, text, level="info": messages.append((title, text, level)))
    monkeypatch.setattr(ui_common, "ask_yes_no", lambda *a, **k: True)
    from app_window import MainWindow
    w = MainWindow()
    w.messages = messages
    w.show()
    yield w
    w.close()


def pump(qapp, until=None, timeout=20.0):
    end = time.time() + timeout
    while time.time() < end:
        qapp.processEvents()
        if until is not None and until():
            return True
        time.sleep(0.01)
    return until is None


def _ftir_trace():
    x = np.linspace(4000, 400, 1801)
    y = 0.02 + 0.8 * np.exp(-4 * np.log(2) * ((x - 1720) / 14) ** 2) + 0.5 * np.exp(-4 * np.log(2) * ((x - 2920) / 18) ** 2)
    t = Trace("a.csv", x, y, "#0072B2")
    import ftir_analysis
    t.peaks = ftir_analysis.detect_peaks(x, y, prominence_frac=0.02)
    t.fits, _ = peak_fitting.fit_peaks_batch(x, y, t.peaks, 25)
    return t


# ------------------------------------------------------------------ exports (no display)
def test_figure_has_every_visible_trace_and_a_composite_fit():
    t = _ftir_trace()
    hidden = Trace("hidden.csv", t.x, t.y * 0.5, "#D55E00")
    hidden.visible = False
    spec = exports.spectrum_figure("ftir", [t, hidden], t.id)
    assert spec.invert_x and [s.role for s in spec.series] == ["data", "fit"]
    model = spec.series[1].y
    peak = int(np.argmin(abs(t.x - 1720)))
    assert abs(model[peak] - t.y[peak]) < 0.08, "the composite follows the fitted band"
    assert np.isfinite(model).all()
    far = int(np.argmin(abs(t.x - 2400)))
    assert model[far] == pytest.approx(t.y[far]), "away from every fitted peak the model is the data itself"


def test_origin_book_and_labtalk_package(tmp_path):
    from labkit.origin import write_labtalk_package
    t = _ftir_trace()
    books = exports.origin_books("ftir", [t], t.id)
    assert len(books) == 1 and books[0].results
    result = write_labtalk_package(books, tmp_path)
    assert result.graphs == 1
    assert "%%" not in (tmp_path / "AeroLab_build.ogs").read_text(encoding="utf-8")


def test_excel_report_has_a_chart_and_sheet_per_trace(tmp_path):
    import openpyxl
    from labkit.excel import write_workbook
    a, b = _ftir_trace(), _ftir_trace()
    b.label = "b.csv"
    a.matches = [{"name": "Poly(x)", "category": "polymers", "score": 0.5, "matched_count": 1,
                  "total_reference_peaks": 2, "source": "ref", "matches": []}]
    path = write_workbook(exports.excel_report("ftir", [a, b], a.id, version="3.0.0"), tmp_path / "r.xlsx")
    book = openpyxl.load_workbook(path)
    assert len(book["Charts"]._charts) == 2
    assert any(name.startswith("Matches") for name in book.sheetnames)
    assert any(name.startswith("Fits") for name in book.sheetnames)


def test_xrd_results_include_sizes_and_williamson_hall():
    t = Trace("p.xy", [1, 2], [1, 2], "#0072B2")
    t.peaks = [{"two_theta": 28.4, "intensity": 10, "fwhm_deg": 0.2, "d_A": 3.14, "size_nm": 40.0}]
    t.metadata["wh_result"] = {"crystallite_size_nm": 35.0, "microstrain": -0.001, "strain_physical": False,
                               "n_peaks": 4, "slope": 0, "intercept": 0}
    names = {r[0]: r for r in exports.results("xrd", t, wavelength=1.540562)}
    assert names["Mean Scherrer crystallite size"][1] == 40.0
    assert "~0" in names["Williamson–Hall microstrain"][3]


# ------------------------------------------------------------------ window
def test_every_page_builds_in_both_themes_without_shortcut_clashes(window, qapp):
    for theme in ("dark", "light"):
        window.apply_theme(theme)
        for key in window.page_keys:
            window.go_to(key)
            pump(qapp, timeout=0.05)
            assert window.is_page_built(key)
    assert window.shortcut_conflicts() == []


def test_ftir_load_detect_undo_redo(window, qapp, ftir_csv):
    ftir = window.page("ftir")
    assert ftir.load_paths([ftir_csv]) == [ftir_csv]
    ftir.detect_peaks()
    t = ftir.get_active()
    assert len(t.peaks) >= 4 and ftir.peaks_table.rowCount() == len(t.peaks)
    window.undo()
    assert not ftir.get_active().peaks
    window.undo()
    assert ftir.traces == [] and ftir.peaks_table.rowCount() == 0
    window.redo()
    window.redo()
    assert len(ftir.get_active().peaks) >= 4


def test_ftir_matching_and_fitting_run_in_the_background(window, qapp, ftir_csv):
    ftir = window.page("ftir")
    ftir.load_paths([ftir_csv])
    ftir.detect_peaks()
    ftir.match_database()
    assert ftir.is_busy
    assert pump(qapp, lambda: not ftir.is_busy)
    t = ftir.get_active()
    assert ftir.matches_table.rowCount() == len(ftir.visible_matches(t))
    ftir.analyse_mixture()
    assert pump(qapp, lambda: not ftir.is_busy)
    assert "mixture" in t.metadata
    ftir.fit_peaks()
    assert pump(qapp, lambda: not ftir.is_busy)
    assert len(t.fits) >= 3 and ftir.fits_table.rowCount() == len(t.fits)
    assert [s.role for s in ftir.figure_spec().series][-1] == "fit"
    window.undo()
    assert not ftir.get_active().fits


def test_processing_asks_then_applies(window, qapp, ftir_csv, monkeypatch):
    import ui_common
    ftir = window.page("ftir")
    ftir.load_paths([ftir_csv])
    before = ftir.get_active().y.copy()
    monkeypatch.setattr(ui_common, "ask_fields", lambda *a, **k: {"wl": "21", "po": "3"})
    ftir.smooth_active()
    assert not np.allclose(before, ftir.get_active().y)
    monkeypatch.setattr(ui_common, "ask_fields", lambda *a, **k: None)
    after = ftir.get_active().y.copy()
    ftir.smooth_active()
    assert np.allclose(after, ftir.get_active().y), "cancelling a dialog changes nothing"
    ftir.revert_active()
    assert np.allclose(before, ftir.get_active().y)


def test_xrd_calculations(window, qapp, xrd_xy, monkeypatch):
    import ui_common
    window.go_to("xrd")
    pump(qapp, timeout=0.1)
    xrd = window.page("xrd")
    xrd.load_paths([xrd_xy])
    xrd.detect_peaks()
    t = xrd.get_active()
    assert len(t.peaks) >= 4
    xrd.run_all_calcs()
    assert abs(t.peaks[0]["d_A"] - 3.1357) < 0.01 and t.peaks[0]["size_nm"] > 0
    xrd.run_williamson_hall()
    assert "Williamson" in xrd.wh_text.toPlainText()
    monkeypatch.setattr(ui_common, "ask_fields", lambda *a, **k: {"crys": "27-30,46-49", "amorph": "12-20"})
    xrd.run_crystallinity()
    assert 0 < t.metadata["crystallinity_pct"] <= 100


def test_custom_wavelength(window, qapp, monkeypatch):
    import ui_common
    xrd = window.page("xrd")
    monkeypatch.setattr(ui_common, "ask_fields", lambda *a, **k: {"wl": "1.789"})
    xrd.wl_combo.setCurrentText("Custom...")
    assert xrd.get_wavelength() == pytest.approx(1.789)


def test_session_round_trip_and_wrong_page(window, qapp, ftir_csv, tmp_path):
    ftir = window.page("ftir")
    ftir.load_paths([ftir_csv])
    ftir.detect_peaks()
    ftir.rb_transmittance.setChecked(True)
    path = ftir.save_session(str(tmp_path / "s.ftirxrd"))
    ftir._on_trace_clear()
    ftir.rb_absorbance.setChecked(True)
    assert ftir.load_session(path)
    assert len(ftir.traces) == 1 and ftir.get_active().peaks and ftir.mode() == "transmittance"
    window.go_to("xrd")
    pump(qapp, timeout=0.1)
    assert window.page("xrd").load_session(path) is False
    assert window.messages[-1][0] == "Wrong page"


def test_reports_and_csv_exports(window, qapp, ftir_csv, tmp_path):
    ftir = window.page("ftir")
    ftir.load_paths([ftir_csv])
    ftir.detect_peaks()
    pdf = ftir.export_pdf_report(str(tmp_path / "r.pdf"))
    assert pdf and Path(pdf).stat().st_size > 5000
    ftir.export_originlab_csv(str(tmp_path / "o.csv"))
    rows = list(csv.reader(open(tmp_path / "o.csv", encoding="utf-8")))
    assert rows[0] == ["sample_ftir.csv X", "sample_ftir.csv Y"] and len(rows) == 2 + len(ftir.get_active().x)


def test_bad_file_is_reported_not_loaded(window, qapp, tmp_path):
    bad = tmp_path / "nonsense.csv"
    bad.write_text("hello\nworld\n", encoding="utf-8")
    ftir = window.page("ftir")
    assert ftir.load_paths([str(bad)]) == []
    assert window.messages and window.messages[-1][0] == "Could not load file"


def test_open_files_routes_by_extension(window):
    assert window.route_for("x.jdx") == "ftir"
    assert window.route_for("x.xrdml") == "xrd"
    window.go_to("xrd")
    assert window.route_for("x.csv") == "xrd"


def test_actions_without_data_warn(window, monkeypatch):
    levels = []
    monkeypatch.setattr(window, "notify", lambda text, level="info": levels.append(level))
    window.page("ftir").detect_peaks()
    window.page("ftir").fit_peaks()
    assert levels == ["warning", "warning"]


def test_database_page_filters(window, qapp):
    page = window.page("database")
    total = page.materials.rowCount()
    assert total > 100
    page.category.setCurrentIndex(page.category.findData("salts_inorganic"))
    assert 0 < page.materials.rowCount() < total
    page.materials.selectRow(0)
    assert page.peaks.rowCount() > 0


# ------------------------------------------------------------------ FTIR screening workflow
PET = "Polyethylene Terephthalate (PET)"


@pytest.fixture
def pet_csv(tmp_path):
    import json
    import ftir_analysis
    import ftir_matching
    db = json.load(open(os.path.join(os.path.dirname(__file__), "..", "database", "ftir_reference_db.json"),
                        encoding="utf-8"))
    x = np.linspace(4000, 600, 1701)
    y = 0.02 + 0.8 * ftir_matching.synthetic_reference_spectrum(ftir_analysis.find_entry(db, PET), x)
    y = y + np.random.default_rng(3).normal(0, 0.004, len(x))
    path = tmp_path / "pet.csv"
    path.write_text("wavenumber,absorbance\n" + "".join(f"{a:.2f},{b:.5f}\n" for a, b in zip(x, y)))
    return str(path)


def _matched(window, qapp, path):
    ftir = window.page("ftir")
    ftir.load_paths([path])
    ftir.detect_peaks()
    ftir.match_database()
    assert pump(qapp, lambda: not ftir.is_busy)
    return ftir, ftir.get_active()


def test_pet_spectrum_is_screened_with_evidence_and_parameters(window, qapp, pet_csv):
    ftir, t = _matched(window, qapp, pet_csv)
    assert t.matches[0]["name"] == PET and t.matches[0]["tier"] == "Strong"
    assert ftir.matches_table.rowCount() == len([m for m in t.matches if m["confidence"] >= ftir.min_conf_spin.value()])
    ftir.matches_table.selectRow(0)
    text = ftir.match_detail.toPlainText()
    assert PET in text and "screening" in text and "✓" in text
    assert ftir.reference_overlay and ftir.reference_overlay["bands"]
    params = t.metadata["match_params"]
    assert params["tolerance_cm1"] == ftir.tolerance_spin.value() and params["materials_searched"] > 200
    ftir.match_filter.setText("terephthal")
    assert all("terephthal" in m["name"].lower() for m in ftir.visible_matches())


def test_processing_or_redetection_clears_stale_matches(window, qapp, pet_csv, monkeypatch):
    import ui_common
    ftir, t = _matched(window, qapp, pet_csv)
    ftir.matches_table.selectRow(0)
    monkeypatch.setattr(ui_common, "ask_fields", lambda *a, **k: {"wl": "11", "po": "3"})
    ftir.smooth_active()
    t = ftir.get_active()
    assert t.matches == [] and t.fg_hits == [] and "match_params" not in t.metadata
    assert ftir.reference_overlay is None and ftir._current_match is None
    assert t.metadata["processing"][-1].startswith("Smooth")
    ftir.detect_peaks()
    ftir.match_database()
    assert pump(qapp, lambda: not ftir.is_busy)
    ftir.detect_peaks()
    assert ftir.get_active().matches == []


def test_undo_while_matching_discards_the_result(window, qapp, pet_csv):
    ftir = window.page("ftir")
    ftir.load_paths([pet_csv])
    ftir.detect_peaks()
    before = window.history.undo_label() if hasattr(window.history, "undo_label") else None
    ftir.analyse_mixture()
    window.undo()                                   # back to "load"
    assert pump(qapp, lambda: not ftir.is_busy)
    assert "mixture" not in ftir.get_active().metadata
    if before is not None:
        assert "mixture" not in (window.history.undo_label() or "")


def test_verdict_session_round_trip_and_exports(window, qapp, pet_csv, tmp_path, monkeypatch):
    import openpyxl
    import ui_common
    from labkit.excel import write_workbook
    import exports
    ftir, t = _matched(window, qapp, pet_csv)
    ftir.matches_table.selectRow(0)
    monkeypatch.setattr(ui_common, "ask_fields", lambda *a, **k: {"note": "SDBS no. 1234"})
    ftir.set_verdict("confirmed")
    assert t.metadata["verdicts"][PET]["status"] == "confirmed"
    assert ftir.matches_table.cell_text(0, 7) == "confirmed"

    pdf = ftir.export_pdf_report(str(tmp_path / "r.pdf"))
    assert pdf and os.path.getsize(pdf) > 10_000
    book = openpyxl.load_workbook(write_workbook(exports.excel_report("ftir", [t], t.id), tmp_path / "r.xlsx"))
    assert any(n.startswith("Match settings") for n in book.sheetnames)

    session = ftir.save_session(str(tmp_path / "s.ftirxrd"))
    assert os.path.getsize(session) < 1_500_000
    assert ftir.load_session(session)
    t2 = ftir.get_active()
    assert t2.matches[0]["name"] == PET and t2.metadata["verdicts"][PET]["note"] == "SDBS no. 1234"
    low = t2.matches[-1]                            # summary-only result: detail rebuilt on demand
    ftir.match_filter.setText(low["name"][:12])
    ftir._show_match_detail(low)
    assert low["name"] in ftir.match_detail.toPlainText()


def test_legacy_coverage_results_are_listed_without_scores(window, qapp, pet_csv):
    ftir = window.page("ftir")
    ftir.method_combo.setCurrentIndex(ftir.method_combo.findData("coverage"))
    assert not ftir.min_conf_spin.isEnabled()
    ftir, t = _matched(window, qapp, pet_csv)
    assert t.matches[0]["method"] == "coverage"
    assert ftir.matches_table.rowCount() == len(t.matches)
    assert "coverage" in ftir.matches_table.cell_text(0, 2)


# ------------------------------------------------------------------ XRD workflow
@pytest.fixture
def ceo2_xy(tmp_path):
    """CeO2-like pattern, 80 nm crystallites, 0.08 deg Gaussian instrument width, with Ka2."""
    l1, l2, a, K, D, inst = 1.540562, 1.544390, 5.411, 0.9, 80.0, 0.08
    x = np.arange(20, 100, 0.01)
    y = np.full_like(x, 50.0)
    for hkl in [(1, 1, 1), (2, 0, 0), (2, 2, 0), (3, 1, 1), (2, 2, 2), (4, 0, 0), (3, 3, 1), (4, 2, 0)]:
        d = a / np.sqrt(sum(v * v for v in hkl))
        for lam, w in ((l1, 1.0), (l2, 0.5)):
            tt = 2 * np.degrees(np.arcsin(lam / (2 * d)))
            f = np.sqrt(np.degrees(K * lam * 0.1 / (D * np.cos(np.radians(tt / 2)))) ** 2 + inst ** 2)
            y += 1000 * w * (0.5 / (1 + 4 * ((x - tt) / f) ** 2) + 0.5 * np.exp(-4 * np.log(2) * ((x - tt) / f) ** 2))
    y = np.random.default_rng(0).poisson(y).astype(float)
    path = tmp_path / "CeO2.xy"
    path.write_text("".join(f"{a:.3f} {b:.0f}\n" for a, b in zip(x, y)))
    return str(path)


def test_xrd_kalpha2_instrument_and_williamson_hall_recover_the_size(window, qapp, ceo2_xy, monkeypatch):
    import ui_common
    xrd = window.page("xrd")
    window.go_to("xrd")
    xrd.load_paths([ceo2_xy])
    answers = iter([{"w": "2.0"}, {"l1": "1.540562", "l2": "1.544390", "r": "0.5"},
                    {"const": "0.08", "U": "0", "V": "0", "W": "0", "mode": "gaussian"}])
    monkeypatch.setattr(ui_common, "ask_fields", lambda *a, **k: next(answers))
    xrd.subtract_background_active()
    xrd.strip_kalpha2_active()
    xrd.edit_instrument_profile()
    t = xrd.get_active()
    assert any("Kα2 stripped" in step for step in t.metadata["processing"])
    xrd.detect_peaks()
    assert 8 <= len(t.peaks) <= 14
    xrd.fit_peaks()
    assert pump(qapp, lambda: not xrd.is_busy)
    xrd.run_all_calcs()
    sizes = [p["size_nm"] for p in t.peaks if p.get("size_nm")]
    assert len(sizes) >= 6 and 65 < np.median(sizes) < 95
    xrd.run_williamson_hall()
    wh = t.metadata["wh_result"]
    assert 60 < wh["crystallite_size_nm"] < 100 and wh["crystallite_size_se_nm"] is not None
    assert t.metadata["xrd_params"]["instrument"]["source"] == "entered"
    assert "Williamson" in xrd.wh_text.toPlainText()


def test_xrd_wavelength_change_clears_stale_results(window, qapp, xrd_xy):
    xrd = window.page("xrd")
    window.go_to("xrd")
    xrd.load_paths([xrd_xy])
    xrd.run_all_calcs()
    t = xrd.get_active()
    assert t.peaks[0]["d_A"]
    xrd.wl_combo.setCurrentText("Mo Ka1")
    assert not t.peaks[0].get("d_A") and "xrd_params" not in t.metadata
    xrd.run_all_calcs()
    assert t.metadata["xrd_params"]["wavelength_name"] == "Mo Ka1"


def test_xrd_custom_wavelength_cancel_reverts(window, qapp, monkeypatch):
    import ui_common
    xrd = window.page("xrd")
    monkeypatch.setattr(ui_common, "ask_fields", lambda *a, **k: None)
    xrd.wl_combo.setCurrentText("Custom...")
    assert xrd.wl_combo.currentText() == "Cu Ka1"


def test_xrd_crystallinity_uses_raw_data_and_peak_toggle(window, qapp, xrd_xy, monkeypatch, tmp_path):
    import ui_common
    import exports
    import openpyxl
    from labkit.excel import write_workbook
    xrd = window.page("xrd")
    window.go_to("xrd")
    xrd.load_paths([xrd_xy])
    answers = iter([{"crys": "27-30,46-49", "amorph": "12-20"}, {"w": "2"}, {"crys": "27-30,46-49", "amorph": "12-20"}])
    monkeypatch.setattr(ui_common, "ask_fields", lambda *a, **k: next(answers))
    xrd.run_crystallinity()
    t = xrd.get_active()
    before = t.metadata["crystallinity_pct"]
    xrd.subtract_background_active()
    xrd.run_crystallinity()
    assert t.metadata["crystallinity_pct"] == pytest.approx(before)
    xrd.detect_peaks()
    xrd.run_all_calcs()
    item = xrd.peaks_table.item(0, 0)
    item.setCheckState(Qt.Unchecked)
    assert t.peaks[0]["use"] is False
    xrd.run_all_calcs()
    assert t.peaks[0].get("size_nm") is None and t.peaks[0]["note"] == "excluded by you"
    pdf = xrd.export_pdf_report(str(tmp_path / "x.pdf"))
    assert pdf and os.path.getsize(pdf) > 10_000
    book = openpyxl.load_workbook(write_workbook(xrd.excel_report(), tmp_path / "x.xlsx"))
    assert book.sheetnames


def test_xrd_lattice_parameter_needs_whole_hkl(window, qapp, xrd_xy, monkeypatch):
    import ui_common
    xrd = window.page("xrd")
    xrd.load_paths([xrd_xy])
    monkeypatch.setattr(ui_common, "ask_fields", lambda *a, **k: {"tt": "28.44", "h": "1.5", "k": "1", "l": "1"})
    xrd.run_lattice_param()
    assert window.messages[-1][0] == "Lattice parameter failed"
    monkeypatch.setattr(ui_common, "ask_fields", lambda *a, **k: {"tt": "28.44", "h": "1", "k": "1", "l": "1"})
    xrd.run_lattice_param()
    assert "a = 5.43" in xrd.calc_log[-1]["result"] and "face-centred" in xrd.calc_log[-1]["result"]
